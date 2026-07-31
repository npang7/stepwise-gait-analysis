from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("StepWise_步态数据分析数学依据说明.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table, header=True):
    table.style = "Table Grid"
    table.autofit = True
    for row_i, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                    run.font.size = Pt(9.5)
            if header and row_i == 0:
                set_cell_shading(cell, "EAF2F8")


def add_para(doc, text="", style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5)


def add_formula_block(doc, title, formula, variables, source, project_use, limitation):
    doc.add_heading(title, level=2)
    table = doc.add_table(rows=6, cols=2)
    labels = ["公式", "变量含义", "依据来源", "本项目用法", "工程阈值/判据", "限制说明"]
    values = [formula, variables, source, project_use, "只输出 screening indicator / candidate，不输出临床诊断。", limitation]
    for i, (label, value) in enumerate(zip(labels, values)):
        set_cell_text(table.cell(i, 0), label, bold=True)
        set_cell_text(table.cell(i, 1), value)
        set_cell_shading(table.cell(i, 0), "F4F6F8")
    style_table(table, header=False)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for name, size, color in [
        ("Title", 22, RGBColor(23, 50, 77)),
        ("Heading 1", 15, RGBColor(23, 50, 77)),
        ("Heading 2", 12.5, RGBColor(31, 70, 105)),
    ]:
        styles[name].font.name = "Arial"
        styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = color

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("StepWise 步态数据分析数学依据说明")
    r.bold = True
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("压力传感器 + 足部 IMU 的工程筛查指标、公式、来源与限制").italic = True

    add_para(
        doc,
        "本文档说明 StepWise 当前 Python 数据分析脚本中每一个主要结论或指标的数学依据。"
        "本系统定位为 portable gait screening and rehabilitation-support tool，不是临床诊断设备。"
    )

    doc.add_heading("1. 结论边界", level=1)
    add_bullet(doc, "可以严谨输出：触地/离地分段、单脚 stance time、stride time 候选、足底前后/内外压力分布、push-off 候选、足部姿态变化。")
    add_bullet(doc, "不应直接输出：扁平足、足外翻、帕金森、脑卒中后步态等医学诊断。最多写 abnormal-like pattern / candidate。")
    add_bullet(doc, "当前示例数据真实采样率约 3.27 Hz，仅适合验证解析与可视化流程；严谨事件检测建议稳定 50-100 Hz。")

    doc.add_heading("2. 传感器布置依据", level=1)
    table = doc.add_table(rows=1, cols=5)
    headers = ["编号", "推荐位置", "英文解剖/工程名称", "用于判断", "依据"]
    for i, h in enumerate(headers):
        set_cell_text(table.cell(0, i), h, bold=True)
    rows = [
        ("P1", "足跟中心或略偏后", "heel / calcaneus", "initial contact, stance start", "Pappas 2004; Sensors 2019"),
        ("P2", "前掌内侧，第一跖骨头", "1st metatarsal head", "medial forefoot load", "Sensors 2019; Pappas 2004"),
        ("P3", "前掌外侧，第五跖骨头", "5th metatarsal head", "lateral forefoot load", "Sensors 2019; Pappas 2004"),
        ("P4", "大脚趾下方", "hallux / great toe", "toe-off, push-off candidate", "Sensors 2019; Sensors 2020"),
        ("IMU", "足背中央或鞋面硬质固定处", "foot-worn IMU", "orientation, angular velocity, event辅助", "gaitmap; Pappas 2004"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val)
    style_table(table)

    add_para(
        doc,
        "如果硬件允许 5 个压力传感器，建议增加 3rd metatarsal head。文献中 5-FSR 鞋垫常用 heel、"
        "1st/3rd/5th metatarsal heads、great toe；3-FSR 标注研究也覆盖大脚趾、跖趾关节、跟骨区。"
    )

    doc.add_heading("3. 数据预处理公式", level=1)
    add_formula_block(
        doc,
        "3.1 真实时间轴与采样率",
        "t_i = timestamp_i - timestamp_0;  fs_est = 1 / median(t_i - t_{i-1})",
        "t_i 为第 i 个采样点相对起始时刻的秒数；fs_est 为由时间戳估计的实际采样率。",
        "gaitmap 开源生态强调 IMU gait analysis 依赖标准化数据结构、采样率与事件时间；步态时空参数均基于时间事件定义。",
        "脚本不用文件头的 100 Hz 直接计算时长，而是用 SystemTime 计算 Time_s、duration、stance/stride time。",
        "如果 ESP32 打包/发送导致时间戳不稳定，必须先修正采样链路；否则所有时间参数都会受影响。",
    )
    add_formula_block(
        doc,
        "3.2 压力非负化与平滑",
        "P'_k(t) = mean(median(max(P_k(t), 0), window), window)",
        "P_k 为第 k 个压力通道；max 保证负压力不进入分析；median/mean rolling 用于降低毛刺。",
        "FSR 鞋垫研究通常先从单通道或多通道压力信号得到稳定的 stance/step pressure peak；实际工程中需抑制毛刺。",
        "脚本对 P1-P4 分别生成 P1_smooth ... P4_smooth，然后所有压力比例和触地判断都基于平滑值。",
        "平滑窗口会引入小的时间延迟；正式实验中应报告窗口大小，并用视频或人工标注验证误差。",
    )

    doc.add_heading("4. 触地/步态周期公式", level=1)
    add_formula_block(
        doc,
        "4.1 总压力",
        "P_total(t) = Σ P'_k(t), k ∈ {1,2,3,4}",
        "P_total 是四个压力点的合力近似。",
        "Sensors 2019 讨论了多 FSR 的 average/sum/cumulative sum 用于 stance 与 step detection；Pappas 2004 使用 heel 与 metatarsal FSR 判断 gait phases。",
        "脚本用 TotalPressure 做 foot contact 的主判据。",
        "四个 FSR 不能得到完整足底压力云图，只能得到低分辨率 plantar loading proxy。",
    )
    add_formula_block(
        doc,
        "4.2 自适应触地阈值",
        "θ_enter = max(θ_min, r · max(P_total));  θ_exit = 0.55 · θ_enter",
        "θ_enter 为进入触地阈值；θ_exit 为退出阈值；r 默认为 0.08；θ_min 默认为 5 N。",
        "FSR 研究常用相对最大压力百分比设阈值，例如 Sensors 2019 对 heel FSR 使用 50% of maximum pressure 作为 stance 条件之一。",
        "脚本使用低比例自适应阈值来适配不同体重/鞋垫标定状态，并用退出阈值形成 hysteresis。",
        "阈值是工程参数，不是医学阈值；正式版本应通过人工视频标注或参考系统调参。",
    )
    add_formula_block(
        doc,
        "4.3 滞回触地状态",
        "contact_i = 1 if P_total(i) ≥ θ_enter; contact_i = 0 if P_total(i) ≤ θ_exit; otherwise keep previous state",
        "contact_i 表示第 i 个采样点是否处于 foot-ground contact / stance。",
        "步态定义中 stance phase 是脚接触地面的阶段，swing phase 是脚离地摆动阶段；智能鞋综述和步态综述均采用此类定义。",
        "脚本通过 contact 状态从 0→1 得到初始触地候选，从 1→0 得到离地候选。",
        "单脚鞋垫只能得到该脚的 stance/swing 候选，不能计算双支撑时间或左右对称性。",
    )
    add_formula_block(
        doc,
        "4.4 Stance time 与 Stride time",
        "T_stance,j = t_end,j - t_start,j;  T_stride,j = t_start,j - t_start,j-1;  T_swing,j ≈ T_stride,j - T_stance,j",
        "j 为第 j 个触地段；start/end 来自 contact 段边界。",
        "步态周期定义为同一只脚连续两次 initial contact 之间；综述文献给出 stance/swing 是基本时空参数。",
        "脚本输出 StanceTime_s、StrideTime_s、SwingTime_s。",
        "只有右脚单鞋垫时，cadence/stride 是单脚估计；需要双脚鞋垫才能得到左右 step time、double support 和 asymmetry。",
    )

    doc.add_heading("5. 压力分布与候选异常公式", level=1)
    add_formula_block(
        doc,
        "5.1 后足/前足压力比例",
        "R_rear = P_rear / P_total;  R_front = P_front / P_total",
        "P_rear = Σ heel channels；P_front = Σ metatarsal + toe channels。",
        "Plantar pressure distribution 是鞋垫步态分析的核心；FSR 位置覆盖 heel、metatarsal、toe 是为了描述触地到推蹬的压力转移。",
        "脚本用 RearRatio_mean 与 FrontRatio_mean 判断 rearfoot-heavy 或 forefoot-heavy candidate。",
        "阈值 0.70 是工程筛查阈值；必须用你们自己的健康样本建立 baseline 后再固定。",
    )
    add_formula_block(
        doc,
        "5.2 内侧/外侧压力比例",
        "R_medial = P_medial / P_total;  R_lateral = P_lateral / P_total",
        "P_medial 通常来自 1st metatarsal 区；P_lateral 通常来自 5th metatarsal 区。",
        "Pappas 2004 在两个跖骨头下放 FSR 是因为足部可能不对称受力；Sensors 2019 也使用 1st/3rd/5th metatarsal heads。",
        "脚本用 MedialRatio_mean 与 LateralRatio_mean 提示 medial/lateral overload candidate。",
        "四点压力无法严谨诊断足内翻/外翻，只能提示偏载；需要更多压力点或临床足底压力板验证。",
    )
    add_formula_block(
        doc,
        "5.3 Late stance 脚趾推蹬比例",
        "R_toe,late = mean(P_toe / P_total), t ∈ last 35% of stance",
        "P_toe 为 hallux/great toe 通道；late stance 取单次 stance 后 35% 时间窗口。",
        "步态接触阶段包括 heel strike、full contact、heel off、toe off；toe 区压力用于观察 terminal stance / toe-off 推蹬。",
        "脚本若 ToeRatio_late_stance_mean 偏低，输出 weak push-off candidate。",
        "如果 P4 未放在大脚趾下方，这个指标不能解释为 push-off；必须先固定传感器布置。",
    )
    add_formula_block(
        doc,
        "5.4 压力冲量",
        "Impulse_j = ∫ P_total(t) dt ≈ trapezoid(P_total, t), t ∈ stance_j",
        "压力冲量表示一次触地期间压力随时间累积的量，单位 N·s。",
        "压力峰值和 stance 期间累计压力常用于 FSR step/gait 研究；Sensors 2019 使用 stance 期间 cumulative sum 降低双峰误计数。",
        "脚本输出 PressureImpulse_Ns，用于比较每一步承重强度。",
        "未经标定的 FSR 非线性很强；压力冲量只有在标定后才有物理量纲意义。",
    )

    doc.add_heading("6. IMU 指标公式", level=1)
    add_formula_block(
        doc,
        "6.1 加速度模长",
        "|a| = sqrt(AccX^2 + AccY^2 + AccZ^2)",
        "AccX/Y/Z 为 IMU 三轴加速度。",
        "IMU gait analysis 常使用三轴加速度、角速度及其特征进行 gait detection、event detection 和 parameter calculation；gaitmap 提供相关开源流程。",
        "脚本输出 AccMag 与每步 AccMag_peak，作为足部运动强度和冲击候选指标。",
        "加速度包含重力分量；如需严谨运动加速度，应进行姿态解算和重力去除。",
    )
    add_formula_block(
        doc,
        "6.2 角速度模长",
        "|ω| = sqrt(GyrX^2 + GyrY^2 + GyrZ^2)",
        "GyrX/Y/Z 为三轴角速度。",
        "Pappas 2004 使用鞋垫中的 gyroscope 结合 FSR 判断 stance、heel-off、swing、heel-strike；gaitmap 也包含 foot-worn IMU event detection。",
        "脚本输出 GyrMag 与每步 GyrMag_peak，用于辅助观察足部旋转强度。",
        "若 IMU 固定不牢，角速度会包含鞋面晃动/线缆扰动，不应直接解释为关节运动。",
    )
    add_formula_block(
        doc,
        "6.3 姿态角范围",
        "Range_pitch,j = max(Pitch_j) - min(Pitch_j); Range_roll,j = max(Roll_j) - min(Roll_j)",
        "j 为单个 stance 段。",
        "足部 IMU 的姿态角和角速度常用于 gait phase / event detection；gaitmap 提供方向对齐、事件检测和参数计算模块。",
        "脚本输出 PitchRange_deg、RollRange_deg，用于筛查每步足部姿态变化幅度。",
        "Yaw 易受磁干扰；MPU6050 无磁力计时 yaw 长时间漂移明显。Pitch/Roll 也需统一坐标轴定义。",
    )

    doc.add_heading("7. 输出标签的数学规则", level=1)
    table = doc.add_table(rows=1, cols=4)
    for i, h in enumerate(["标签", "脚本规则", "含义", "严谨写法"]):
        set_cell_text(table.cell(0, i), h, bold=True)
    rows = [
        ("Rearfoot-heavy", "RearRatio_mean ≥ 0.70", "该步后足压力占比高", "rearfoot-heavy contact candidate"),
        ("Forefoot-heavy", "FrontRatio_mean ≥ 0.70", "该步前足压力占比高", "forefoot-heavy contact candidate"),
        ("Medial overload", "MedialRatio_mean ≥ 0.65", "内侧前掌偏载", "medial loading bias candidate"),
        ("Lateral overload", "LateralRatio_mean ≥ 0.65", "外侧前掌偏载", "lateral loading bias candidate"),
        ("Weak push-off candidate", "ToeRatio_late_stance_mean < 0.12", "晚期触地脚趾参与不足", "weak push-off candidate"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val)
    style_table(table)
    add_para(doc, "以上阈值用于 demo 和工程筛查。毕业设计中应说明：阈值来自工程假设，需要通过健康受试者 baseline、视频标注或参考压力平台验证。")

    doc.add_heading("8. 推荐实验与验证方法", level=1)
    add_bullet(doc, "采样率：ESP32 端真实输出应稳定在 50-100 Hz；PC 端保存每包 timestamp，并检查丢包。")
    add_bullet(doc, "压力标定：对每个 FSR 做空载零点、已知砝码加载、多点拟合，保存 raw-to-N calibration curve。")
    add_bullet(doc, "传感器位置：测试前拍照记录 P1-P4 位置；代码配置必须与物理位置一致。")
    add_bullet(doc, "人工标注：同步录制侧面视频，人工标注 initial contact / toe-off，与算法事件比较误差。")
    add_bullet(doc, "重复性：同一人至少 3 次 walking trial，每次不少于 10-20 个有效 stance 段。")
    add_bullet(doc, "报告表述：输出 abnormal-like indicator，不写 clinical diagnosis。")

    doc.add_heading("9. 参考文献与开源来源", level=1)
    refs = [
        "Küderle A. et al. Gaitmap—An Open Ecosystem for IMU-Based Human Gait Analysis and Algorithm Benchmarking. IEEE Open Journal of Engineering in Medicine and Biology, 2024. DOI: 10.1109/OJEMB.2024.3356791. https://pmc.ncbi.nlm.nih.gov/articles/PMC10939318/",
        "gaitmap documentation: event detection, stride segmentation, trajectory reconstruction, parameter calculation. https://gaitmap.readthedocs.io/en/stable/modules/event_detection.html",
        "Pappas I. P. I. et al. A Reliable Gyroscope-Based Gait-Phase Detection Sensor Embedded in a Shoe Insole. IEEE Sensors Journal, 2004, 4(2):268-274. DOI: 10.1109/JSEN.2004.823671. https://scienceportal.tecnalia.com/en/publications/a-reliable-gyroscope-based-gait-phase-detection-sensor-embedded-i-2/",
        "Jeon H. et al. Fast Wearable Sensor-Based Foot-Ground Contact Phase Classification Using a CNN with Sliding-Window Label Overlapping. Sensors, 2020. https://pmc.ncbi.nlm.nih.gov/articles/PMC7506746/",
        "Design and Accuracy of an Instrumented Insole Using Pressure Sensors for Step Count. Sensors, 2019. https://www.mdpi.com/1424-8220/19/5/984",
        "Enhancing Intelligent Shoes with Gait Analysis: A Review on the Spatiotemporal Estimation Techniques. 2024. https://pmc.ncbi.nlm.nih.gov/articles/PMC11678955/",
    ]
    for ref in refs:
        add_bullet(doc, ref)

    doc.add_heading("10. 可直接写入报告的总结句", level=1)
    add_para(
        doc,
        "StepWise extracts engineering gait-screening indicators from low-cost insole pressure sensors and a foot-worn IMU. "
        "The mathematical pipeline follows common wearable gait-analysis practice: timestamp-based temporal parameters, "
        "FSR-based stance/swing segmentation, plantar pressure distribution ratios, pressure impulse, and IMU magnitude/range features. "
        "The output is intended to flag abnormal-like loading and timing patterns, not to provide clinical diagnosis.",
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
